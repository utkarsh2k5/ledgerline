import re
from datetime import date
from pathlib import Path

import pytest
from docx import Document as Docx
from pptx import Presentation
from sqlalchemy import select

from app.analytics.engine import build_analysis
from app.analytics.facts import fmt
from app.config import settings
from app.ingestion.loader import load_file
from app.models import LineageFact, ReportRun
from app.reporting.pipeline import GateError, compose, generate, reproduce
from app.reporting.spec import SpecError, load_spec, parse_spec
from tests.conftest import resolve_all_breaks

PE = date(2026, 9, 30)


def _run(db, scheme, spec="quarterly_trustee_report", override=False):
    run = ReportRun(scheme_id=scheme.id, spec_name=spec, period_end=PE, created_by="test", gate_overridden=override)
    db.add(run)
    db.commit()
    return run


def test_gate_blocks_report_with_open_material_breaks(db, scheme):
    with pytest.raises(GateError, match="5 material"):
        generate(db, _run(db, scheme))


def test_full_report_generation(db, scheme):
    resolve_all_breaks(db, scheme)
    run = generate(db, _run(db, scheme))
    assert run.status == "done" and run.commentary_status == "verified"
    paths = {k: settings.storage_dir / v for k, v in run.outputs.items()}
    assert set(paths) == {"docx", "pptx", "html"}

    # Word: the KPI values shown must be the computed facts
    analysis = build_analysis(db, scheme, PE)
    doc_text = "\n".join(c.text for t in Docx(paths["docx"]).tables for row in t.rows for c in row.cells)
    for fid in ("mv.TOTAL.end", "perf.TOTAL.3m.ret", "funding.ratio"):
        f = analysis.facts[fid]
        assert fmt(f.value, f.unit, "GBP") in doc_text

    # PowerPoint opens and has a title slide plus content
    assert len(Presentation(paths["pptx"]).slides) >= 12

    # HTML embeds clickable facts with lineage
    html = paths["html"].read_text()
    assert 'data-fact="perf.TOTAL.3m.ret"' in html
    assert '"ret.GEQ.2026-09"' in html  # leaf facts embedded for tracing

    # Every displayed fact is stored with its lineage
    stored = {f.fact_id: f for f in db.scalars(select(LineageFact).where(LineageFact.report_run_id == run.id))}
    leaf = stored["ret.GEQ.2026-09"]
    assert leaf.sources and leaf.sources[0].startswith("custodian_valuations")
    assert "Modified Dietz" in leaf.formula


def test_report_is_reproducible_and_detects_data_changes(db, scheme, sample_dir):
    resolve_all_breaks(db, scheme)
    run = generate(db, _run(db, scheme, "trustee_summary"))
    again = reproduce(db, run)
    assert again["identical"] and again["n_changed_facts"] == 0 and again["commentary_still_valid"]

    # A benchmark provider restates one month -> the old report is no longer reproducible
    csv = (sample_dir / "benchmark_returns.csv").read_text().splitlines()
    idx = next(i for i, line in enumerate(csv) if line.startswith("GEQ,2026-09-30"))
    code, month, val = csv[idx].split(",")
    csv[idx] = f"{code},{month},{float(val) + 0.5:.4f}"
    load_file(db, scheme, "benchmark_csv", "benchmark_restated.csv", ("\n".join(csv) + "\n").encode(), "test")
    db.commit()
    after = reproduce(db, run)
    assert not after["identical"]
    changed = {c["fact_id"] for c in after["changed_facts"]}
    assert "perf.TOTAL.3m.bench" in changed
    assert after["commentary_still_valid"] is False  # stored prose quotes the old benchmark figure


def test_gate_override_adds_notice(db, scheme):
    _, doc, _ = compose(db, scheme, load_spec("trustee_summary"), PE, gate_overridden=True)
    assert doc.sections[0].title == "Important notice"


def test_specs_are_validated():
    with pytest.raises(SpecError):
        parse_spec("name: x\ntitle: y\nsections:\n  - type: kpis\n")  # kpis without items
    with pytest.raises(SpecError):
        parse_spec("name: x\ntitle: y\nsections:\n  - type: not_a_section\n")
    with pytest.raises(SpecError):
        load_spec("../../etc/passwd")
